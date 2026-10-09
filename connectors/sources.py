"""Read tabular data from a company's own services. Read-only, size-limited, no secrets."""
import csv
import hashlib
import http.client
import io
import ipaddress
import json
import re
import socket
import ssl
from urllib.parse import parse_qs, urlparse
from urllib.request import Request, build_opener, HTTPRedirectHandler

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLUMNS = 50
MAX_CELL = 500
FETCH_TIMEOUT = 10
SHEET_ID = re.compile(r'^/spreadsheets/d/([A-Za-z0-9_-]{20,100})')

# Universal catalog: what a company can connect. Only 'available' kinds work today.
CATALOG = (
    {'kind': 'csv', 'title': 'Excel / CSV', 'group': 'Таблиці', 'state': 'available',
     'gives': 'Будь-яка таблиця: замовлення, оплати, залишки'},
    {'kind': 'google_sheets', 'title': 'Google Таблиці', 'group': 'Таблиці', 'state': 'available',
     'gives': 'Опублікована таблиця, оновлюється за посиланням'},
    {'kind': 'url', 'title': 'Таблиця за посиланням', 'group': 'Інші сервери', 'state': 'available',
     'gives': 'CSV, JSON або Excel з будь-якого сервера за https-посиланням'},
    {'kind': 'monobank', 'title': 'Monobank', 'group': 'Банк', 'state': 'planned',
     'gives': 'Виписки та надходження'},
    {'kind': 'privatbank', 'title': 'ПриватБанк', 'group': 'Банк', 'state': 'planned',
     'gives': 'Виписки та надходження'},
    {'kind': 'nova_poshta', 'title': 'Нова Пошта', 'group': 'Доставка', 'state': 'planned',
     'gives': 'Статуси відправлень'},
    {'kind': 'checkbox', 'title': 'Checkbox', 'group': 'Каси (ПРРО)', 'state': 'planned',
     'gives': 'Чеки та зміни'},
    {'kind': 'binotel', 'title': 'Binotel', 'group': 'Телефонія', 'state': 'planned',
     'gives': 'Журнал дзвінків'},
    {'kind': 'ringostat', 'title': 'Ringostat', 'group': 'Телефонія', 'state': 'planned',
     'gives': 'Журнал дзвінків'},
    {'kind': 'bas', 'title': 'BAS', 'group': 'Облік', 'state': 'planned',
     'gives': 'Замовлення, залишки, рахунки'},
    {'kind': 'dilovod', 'title': 'Dilovod', 'group': 'Облік', 'state': 'planned',
     'gives': 'Замовлення, залишки, рахунки'},
)


class SourceError(ValueError):
    pass


def _decode(raw):
    for encoding in ('utf-8-sig', 'cp1251'):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise SourceError('Не вдалося прочитати кодування файлу. Збережіть його як CSV UTF-8.')


def _table(rows):
    rows = [[('' if cell is None else str(cell)).strip()[:MAX_CELL] for cell in row] for row in rows]
    rows = [row for row in rows if any(row)]
    if not rows:
        raise SourceError('Таблиця порожня.')
    header, body = list(rows[0]), rows[1:]
    while header and not header[-1]:
        header.pop()
    if len(header) > MAX_COLUMNS:
        raise SourceError(f'Забагато колонок: максимум {MAX_COLUMNS}.')
    if len(body) > MAX_ROWS:
        raise SourceError(f'Забагато рядків: максимум {MAX_ROWS}.')
    columns = [name or f'Колонка {index + 1}' for index, name in enumerate(header)]
    if len(set(columns)) != len(columns):
        raise SourceError('Назви колонок у першому рядку мають бути різними.')
    width = len(columns)
    body = [(row + [''] * width)[:width] for row in body]
    digest = hashlib.sha256(json.dumps([columns, body], ensure_ascii=False).encode()).hexdigest()
    return {'columns': columns, 'rows': body, 'row_count': len(body), 'sha256': digest}


def parse_csv(raw):
    if len(raw) > MAX_BYTES:
        raise SourceError('Файл більший за 5 МБ.')
    text = _decode(raw)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=',;\t')
    except csv.Error:
        dialect = csv.excel
    return _table(csv.reader(io.StringIO(text), dialect))


def parse_xlsx(raw):
    if len(raw) > MAX_BYTES:
        raise SourceError('Файл більший за 5 МБ.')
    from openpyxl import load_workbook
    try:
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise SourceError('Не вдалося відкрити Excel-файл.') from exc
    try:
        sheet = book.worksheets[0]
        return _table(sheet.iter_rows(values_only=True, max_row=MAX_ROWS + 2, max_col=MAX_COLUMNS + 1))
    finally:
        book.close()


def parse_upload(name, raw):
    lower = (name or '').lower()
    if lower.endswith('.xlsx'):
        return parse_xlsx(raw)
    if lower.endswith(('.csv', '.txt')):
        return parse_csv(raw)
    raise SourceError('Підтримуються файли .xlsx та .csv.')


def sheet_export_url(url):
    """Only public Google Sheets on docs.google.com; anything else is refused (no SSRF)."""
    parts = urlparse(url or '')
    match = SHEET_ID.match(parts.path or '')
    if parts.scheme != 'https' or parts.hostname != 'docs.google.com' or parts.port or not match:
        raise SourceError('Потрібне посилання https://docs.google.com/spreadsheets/d/…')
    gid = parse_qs(parts.query).get('gid', [''])[0] or parse_qs(parts.fragment).get('gid', [''])[0]
    gid = gid if gid.isdigit() else '0'
    return f'https://docs.google.com/spreadsheets/d/{match.group(1)}/export?format=csv&gid={gid}'


class _SameHostRedirect(HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        host = urlparse(newurl).hostname or ''
        if urlparse(newurl).scheme != 'https' or not (host == 'docs.google.com'
                                                      or host.endswith('.googleusercontent.com')):
            raise SourceError('Google перенаправив на неочікувану адресу.')
        return super().redirect_request(req, fp, code, msg, headers, newurl)


def fetch_sheet(url, opener=None):
    export = sheet_export_url(url)
    opener = opener or build_opener(_SameHostRedirect)
    try:
        with opener.open(Request(export, headers={'User-Agent': 'BoS-connector'}),
                         timeout=FETCH_TIMEOUT) as response:
            raw = response.read(MAX_BYTES + 1)
    except SourceError:
        raise
    except Exception as exc:
        raise SourceError('Таблиця недоступна. Перевірте, що її опубліковано: Файл → Поділитися → '
                          'Опублікувати в інтернеті → CSV.') from exc
    if raw.lstrip()[:15].lower().startswith(b'<!doctype html') or raw.lstrip()[:5].lower() == b'<html':
        raise SourceError('Таблиця не опублікована або закрита доступом.')
    return parse_csv(raw)


# --- any other server: a public https link to CSV, JSON or Excel ----------------------------------------

JSON_LISTS = ('data', 'items', 'results', 'rows', 'records')


def _cell(value):
    if value is None:
        return ''
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def parse_json(raw):
    """A JSON array of objects (or of rows with a header row), or an object holding one under a common key."""
    if len(raw) > MAX_BYTES:
        raise SourceError('Файл більший за 5 МБ.')
    try:
        data = json.loads(_decode(raw))
    except ValueError as exc:
        raise SourceError('За посиланням некоректний JSON.') from exc
    if isinstance(data, dict):
        data = next((data[key] for key in JSON_LISTS if isinstance(data.get(key), list)), data)
    if not isinstance(data, list) or not data:
        raise SourceError('Очікується список записів JSON: масив об’єктів.')
    if all(isinstance(row, list) for row in data):
        return _table([[_cell(v) for v in row] for row in data])
    if not all(isinstance(row, dict) for row in data):
        raise SourceError('Очікується список записів JSON: масив об’єктів.')
    columns = []
    for row in data:
        for key in row:
            if key not in columns:
                columns.append(key)
        if len(columns) > MAX_COLUMNS:
            raise SourceError(f'Забагато колонок: максимум {MAX_COLUMNS}.')
    return _table([[str(c) for c in columns]] + [[_cell(row.get(c)) for c in columns] for row in data])


def public_address(host, port=443):
    """Resolve the host and refuse anything that is not a public internet address (no SSRF into the LAN)."""
    try:
        infos = socket.getaddrinfo(host, port, type=socket.SOCK_STREAM)
    except OSError as exc:
        raise SourceError('Сервер за посиланням не знайдено.') from exc
    if not infos:
        raise SourceError('Сервер за посиланням не знайдено.')
    for info in infos:
        address = ipaddress.ip_address(info[4][0].split('%')[0])
        if not address.is_global or address.is_multicast:
            raise SourceError('Посилання веде до внутрішньої мережі. Дозволені лише публічні адреси інтернету.')
    return infos[0][4][:2]


def _https_get(host, address, path):
    """One GET to the already-checked address, verified TLS for the host name, no redirects followed."""
    context = ssl.create_default_context()
    raw = socket.create_connection(address, timeout=FETCH_TIMEOUT)
    try:
        sock = context.wrap_socket(raw, server_hostname=host)
    except Exception:
        raw.close()
        raise
    connection = http.client.HTTPSConnection(host, 443, timeout=FETCH_TIMEOUT, context=context)
    connection.sock = sock
    try:
        connection.request('GET', path, headers={'User-Agent': 'BoS-connector',
                                                 'Accept': 'text/csv, application/json;q=0.9, */*;q=0.5'})
        response = connection.getresponse()
        if 300 <= response.status < 400:
            raise SourceError('Посилання перенаправляє на іншу адресу. Вкажіть кінцеве посилання на файл.')
        if response.status != 200:
            raise SourceError(f'Сервер за посиланням відповів кодом {response.status}.')
        return response.read(MAX_BYTES + 1), response.getheader('Content-Type', '') or ''
    finally:
        connection.close()


def fetch_url(url, get=None):
    """Read CSV, JSON or Excel from a public https link of any other server. Read-only, size- and time-limited."""
    parts = urlparse((url or '').strip())
    if parts.scheme != 'https' or not parts.hostname or parts.username or parts.password:
        raise SourceError('Потрібне посилання https://… без логіна й пароля в адресі.')
    try:
        port = parts.port or 443
    except ValueError as exc:
        raise SourceError('Некоректне посилання.') from exc
    if port != 443:
        raise SourceError('Дозволено лише стандартний порт https (443).')
    address = public_address(parts.hostname)
    path = (parts.path or '/') + ('?' + parts.query if parts.query else '')
    try:
        raw, kind = (get or _https_get)(parts.hostname, address, path)
    except SourceError:
        raise
    except Exception as exc:
        raise SourceError('Сервер за посиланням недоступний.') from exc
    head = raw.lstrip()[:15].lower()
    if head.startswith(b'<!doctype html') or head.startswith(b'<html'):
        raise SourceError('За посиланням вебсторінка, а не таблиця. Потрібне пряме посилання на CSV, JSON або Excel.')
    kind = kind.split(';')[0].strip().lower()
    if raw[:2] == b'PK' or kind.endswith('spreadsheetml.sheet'):
        return parse_xlsx(raw)
    if kind.endswith('json') or head[:1] in (b'[', b'{'):
        return parse_json(raw)
    return parse_csv(raw)


def fetch(kind, url):
    """Read a connected source again by its kind; uploaded files are never re-read."""
    if kind == 'google_sheets':
        return fetch_sheet(url)
    if kind == 'url':
        return fetch_url(url)
    raise SourceError('Це джерело не оновлюється за посиланням.')
