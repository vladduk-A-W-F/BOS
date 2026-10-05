import hashlib,io,zipfile
from .models import Document

MAX_BYTES=10*1024*1024

def parse(file):
    try:
        return _parse(file)
    except (ValueError, UnicodeError):
        raise
    except ImportError as exc:
        raise ValueError('Модуль читання документа відсутній. Повторіть встановлення залежностей.') from exc
    except Exception as exc:
        raise ValueError('Не вдалося прочитати документ. Перевірте формат і цілісність файла.') from exc

IMAGE_TYPES={'png':'image/png','jpg':'image/jpeg','jpeg':'image/jpeg'}

def _png_complete(data):
    """Whole PNG chunk chain with valid CRCs: IHDR first, IDAT present, IEND last, nothing after."""
    import zlib
    if data[:8]!=b'\x89PNG\r\n\x1a\n':return False
    pos,first,idat=8,True,False
    while pos+12<=len(data):
        length=int.from_bytes(data[pos:pos+4],'big');kind=data[pos+4:pos+8];end=pos+12+length
        if end>len(data) or zlib.crc32(data[pos+4:pos+8+length])!=int.from_bytes(data[pos+8+length:end],'big'):return False
        if first and (kind!=b'IHDR' or length!=13):return False
        if first and (int.from_bytes(data[16:20],'big')==0 or int.from_bytes(data[20:24],'big')==0):return False
        first=False;idat=idat or kind==b'IDAT';pos=end
        if kind==b'IEND':return idat and length==0 and pos==len(data)
    return False

def _jpeg_complete(data):
    """Marker walk: SOI, sized frame, scans with enough coded bytes for every 8x8 block, final EOI.

    libjpeg silently fills a missing scan with grey, so a decode alone does not prove the photo is whole.
    Lower bound only: each block needs at least 2 coded bits in a baseline scan, 1 bit in a progressive one.
    """
    if data[:2]!=b'\xff\xd8':return False
    pos,blocks,progressive,coded,scan=2,0,False,0,False
    while pos<len(data):
        if data[pos]!=0xFF:return False
        while pos<len(data) and data[pos]==0xFF:pos+=1
        if pos>=len(data):return False
        marker=data[pos];pos+=1
        if marker==0xD9:return blocks>0 and scan and coded*8>=blocks*(1 if progressive else 2) and not data[pos:].strip(b'\x00')
        if marker==0x01 or 0xD0<=marker<=0xD7:continue
        if pos+2>len(data):return False
        length=int.from_bytes(data[pos:pos+2],'big')
        if length<2 or pos+length>len(data):return False
        if marker in (0xC0,0xC1,0xC2,0xC3,0xC5,0xC6,0xC7,0xC9,0xCA,0xCB,0xCD,0xCE,0xCF):
            height,width,count=int.from_bytes(data[pos+3:pos+5],'big'),int.from_bytes(data[pos+5:pos+7],'big'),data[pos+7] if length>7 else 0
            if blocks or length!=8+3*count or not (height and width and count):return False
            factors=[(data[pos+9+3*i]>>4,data[pos+9+3*i]&15) for i in range(count)]
            hmax,vmax=max(h for h,_ in factors),max(v for _,v in factors)
            if not (hmax and vmax):return False
            blocks=sum(-(-(-(-width*h//hmax))//8)*-(-(-(-height*v//vmax))//8) for h,v in factors)
            progressive=marker in (0xC2,0xC6,0xCA,0xCE)
        pos+=length
        if marker==0xDA:
            if not blocks:return False
            scan=True
            # Entropy-coded data runs until a marker that is not stuffing or a restart.
            begin=pos
            while pos+1<len(data) and not (data[pos]==0xFF and data[pos+1] not in (0x00,*range(0xD0,0xD8))):pos+=1
            if pos+1>=len(data):return False
            coded+=pos-begin
    return False

MAX_IMAGE_PIXELS=40_000_000

def _decodes(data,expected):
    """Full decode with Pillow under a pixel limit; any warning or error means not a usable photo."""
    import io,warnings
    from PIL import Image
    try:
        with warnings.catch_warnings():
            warnings.simplefilter('error')
            with Image.open(io.BytesIO(data)) as image:
                if image.format!=expected:return False
                width,height=image.size
                if not (0<width and 0<height and width*height<=MAX_IMAGE_PIXELS):return False
                image.load()
        return True
    except Exception:
        return False

def image_type(name,data):
    """Content type of a PNG/JPEG original that is structurally complete and fully decodes, otherwise None. No OCR."""
    ext=(name or '').rsplit('.',1)[-1].lower()
    if ext=='png' and _png_complete(data) and _decodes(data,'PNG'):return IMAGE_TYPES[ext]
    if ext in ('jpg','jpeg') and _jpeg_complete(data) and _decodes(data,'JPEG'):return IMAGE_TYPES[ext]
    return None

def _parse(file):
    data=file.read(MAX_BYTES+1)
    if len(data)>MAX_BYTES:raise ValueError('Файл перевищує 10 МБ.')
    ext=file.name.rsplit('.',1)[-1].lower();parts=[]
    if ext in ('docx','xlsx'):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(x.file_size for x in z.infolist())>40*1024*1024:raise ValueError('Розпакований документ завеликий.')
    if ext=='pdf':
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(data))
        if reader.is_encrypted:raise ValueError('Зніміть пароль із PDF перед завантаженням.')
        if len(reader.pages)>100:raise ValueError('Максимум 100 сторінок.')
        parts=[{'source':f'Сторінка {i+1}','text':page.extract_text() or ''} for i,page in enumerate(reader.pages)]
    elif ext=='docx':
        from docx import Document as Word
        d=Word(io.BytesIO(data));parts=[{'source':f'Абзац {i+1}','text':p.text} for i,p in enumerate(d.paragraphs) if p.text]
        for i,t in enumerate(d.tables):parts.append({'source':f'Таблиця {i+1}','text':'\n'.join(' | '.join(c.text for c in row.cells) for row in t.rows)})
    elif ext=='xlsx':
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
        for ws in wb:
            if ws.max_row>5000 or ws.max_column>100:raise ValueError('Аркуш перевищує 5000 рядків або 100 стовпців.')
            parts.append({'source':f'Аркуш {ws.title}','text':'\n'.join(' | '.join(str(v) if v is not None else '' for v in row) for row in ws.iter_rows(values_only=True))})
        wb.close()
    elif ext in ('txt','md'):parts=[{'source':'Текст','text':data.decode('utf-8-sig')}]
    elif ext in IMAGE_TYPES:
        # A photo of a report is stored as the verified original; text recognition is not available yet.
        if image_type(file.name,data) is None:raise ValueError('Файл не є цілим зображенням PNG або JPEG.')
    else:raise ValueError('Підтримуються PDF, DOCX, XLSX, TXT, MD, JPEG та PNG.')
    text='\n'.join(p['text'] for p in parts)
    if len(text)>500000:raise ValueError('Текст документа перевищує ліміт.')
    return {'text':text,'sections':parts,'checksum':hashlib.sha256(data).hexdigest(),'status':'needs_review' if text.strip() else 'ocr_required','content':data}
